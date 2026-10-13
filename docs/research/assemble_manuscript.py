"""Assemble the manuscript into the faculty's Word template.

    python docs/research/assemble_manuscript.py            draft: pending markers allowed
    python docs/research/assemble_manuscript.py --final    refuses pending markers and gaps
    powershell -File docs/research/finalize_manuscript.ps1 Word fills page numbers, writes PDF

Needs python-docx and latex2mathml (lxml comes with python-docx), and Office's MML2OMML.XSL,
which turns MathML into native Word equations that open in Word's Equation Editor, as the
guidelines require. The template is the faculty's own .docx in docs/faculty/source/, which is
not committed. Its cover page, Index, List of Figures, List of Tables, Contributions tables, IEEE
header, and the faculty member's biography are kept and filled; its sample body is replaced by
docs/research/00-23. Paragraph formats copy the template's own (faces, sizes, colours, spacing).

Words are counted, for the Contributions table and manuscript_report.json, as the words of prose
paragraphs and list items; headings, tables, captions, equations, and code are not counted.
"""

from __future__ import annotations

import argparse
import copy
import itertools
import json
import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import docx
import latex2mathml.converter
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
RESEARCH = ROOT / "docs" / "research"
TEMPLATE = ROOT / "docs" / "faculty" / "source" / "3_Consolidated_Template.docx"
OUT = ROOT / "docs" / "final" / "TRIAD_Manuscript_26FSPDSCPMID4096.docx"
XSL = Path(os.environ.get("MML2OMML_XSL",
                          r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"))
TEAM_ID = "26FSPDSCPMID4096"
MEMBERS = [f"[TEAM MEMBER {i}]" for i in range(1, 5)]

M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

HELV = "Helvetica Neue"               # the template's heading and caption face
BLUE, GREY = "00629B", "58595B"       # its section and subsection heading colours
BODY, SMALL, CAPTION, HEAD = 20, 16, 14, 18   # half-points: 10, 8, 7, 9 pt
TEXT_TWIPS = 11906 - 2 * 1440         # A4 less the template's 2.54 cm side margins
MAX_FIGURE_IN = 8.5                   # the template's maximum figure depth

FIG = re.compile(r"\*\*Fig\. (\d+)\. (.+?)\*\*\s*\*\((.+?)\)\*$", re.S)
TCAP = re.compile(r"\*\*Table (\d+)\. (.+?)\*\*(?:\s*\*\((.+?)\)\*)?$", re.S)
LIST = re.compile(r"(\s*)([-*]|\d+\.) (.*)")
CITES = r"\[@[^\]\s]+\](?:\s*\[@[^\]\s]+\])*"
TOKEN = re.compile(rf"(?P<code>`[^`]+`)|(?P<math>\$[^$]+\$)|(?P<cite>{CITES})"
                   r"|(?P<bold>\*\*.+?\*\*)|(?P<italic>\*[^*\s][^*]*?\*)")
NUM = r"\d+(?:\.\d+)?"
JOIN = r"(?:,\s*|\s+and\s+|\s+or\s+|\s+to\s+|\s*[-\u2013]\s*)"
SECTION_REF = re.compile(rf"\b(Sections?) ({NUM}(?:{JOIN}{NUM})*)")
TABLE_REF = re.compile(rf"\b(Tables?) (\d+(?:{JOIN}\d+)*)\b")


def roman(n: int) -> str:
    out = ""
    for value, numeral in ((1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"),
                           (90, "XC"), (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"),
                           (4, "IV"), (1, "I")):
        while n >= value:
            out, n = out + numeral, n - value
    return out


def section_label(ref: str) -> str:
    """'17' -> 'XVII', '17.3' -> 'XVII-C': IEEE numbers sections in Roman and letters subsections."""
    major, _, minor = ref.partition(".")
    return roman(int(major)) + (f"-{chr(64 + int(minor))}" if minor else "")


def transform(s: str) -> str:
    """Typesetting for text outside code and maths: numbering, dashes, quotes, number spacing."""
    s = SECTION_REF.sub(lambda m: m[1] + " " + re.sub(NUM, lambda n: section_label(n[0]), m[2]), s)
    s = TABLE_REF.sub(lambda m: m[1] + " " + re.sub(r"\d+", lambda n: roman(int(n[0])), m[2]), s)
    s = re.sub(r"(?<=\S) - (?=\S)", "\u2014", s)                  # the sources' spaced hyphen
    s = re.sub(r"(?<!\d)(\d{1,3}) (?=\d{3}(?!\d))", "\\1\u00a0", s)  # 6 814 stays on one line
    s = re.sub(r'(^|[\s(\[\u2014])"', "\\1\u201c", s).replace('"', "\u201d")
    return re.sub(r"(^|[\s(\[\u2014])'", "\\1\u2018", s).replace("'", "\u2019")


def words(text: str) -> int:
    text = re.sub(CITES, " ", text)
    text = re.sub(r"\$[^$]+\$", " x ", text)
    return sum(1 for w in re.sub(r"[*`]", "", text).split() if re.search(r"\w", w))


# -- reading the sources ---------------------------------------------------------------------

@dataclass
class Block:
    kind: str     # heading, para, list, table, table_caption, figure, equation, code
    text: str = ""
    level: int = 0
    items: list[str] = field(default_factory=list)
    ordered: bool = False
    rows: list[list[str]] = field(default_factory=list)
    number: int = 0
    title: str = ""


def read(path: Path) -> tuple[list[Block], list[str]]:
    """A section's manuscript blocks (its Evidence block dropped) and its pending markers.

    A line `{{include: tables/x.md}}` is replaced by that generated file, so a table built from
    stored results reaches the manuscript as built rather than as last pasted.
    """
    text = path.read_text(encoding="utf-8").replace("\r\n", "\n")
    text = re.split(r"\n## Evidence", text)[0].rstrip()
    text = re.sub(r"\n---$", "", text).rstrip()
    text = re.sub(r"^\{\{include: (.+?)\}\}$",
                  lambda m: (RESEARCH / m[1].strip()).read_text(encoding="utf-8").strip(),
                  text, flags=re.M)
    pending = [c.strip() for c in re.findall(r"<!--(.*?)-->", text, flags=re.S)]
    return parse(re.sub(r"<!--.*?-->", "", text, flags=re.S).split("\n")), pending


def parse(lines: list[str]) -> list[Block]:
    blocks: list[Block] = []
    para: list[str] = []

    def flush() -> None:
        if not para:
            return
        text = " ".join(s.strip() for s in para)
        para.clear()
        if m := FIG.match(text):
            blocks.append(Block("figure", number=int(m[1]), title=m[2].strip()))
        elif m := TCAP.match(text):
            blocks.append(Block("table_caption", number=int(m[1]), title=m[2].strip().rstrip("."),
                                text=m[3] or ""))   # text: a note set under the table
        else:
            blocks.append(Block("para", text=text))

    i = 0
    while i < len(lines):
        line, s = lines[i], lines[i].strip()
        if not s:
            flush()
            i += 1
        elif s.startswith("```"):
            flush()
            j = i + 1
            while not lines[j].strip().startswith("```"):
                j += 1
            blocks.append(Block("code", text="\n".join(lines[i + 1:j])))
            i = j + 1
        elif s == "$$":
            flush()
            j = i + 1
            while lines[j].strip() != "$$":
                j += 1
            blocks.append(Block("equation", text=" ".join(x.strip() for x in lines[i + 1:j])))
            i = j + 1
        elif m := re.match(r"(#{1,3}) (.+)", s):
            flush()
            blocks.append(Block("heading", text=m[2].strip(), level=len(m[1])))
            i += 1
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(Block("table", rows=rows))
        elif not para and (m := LIST.match(line)) and not m[1]:
            ordered, items = m[2][0].isdigit(), []
            while i < len(lines):
                m2 = LIST.match(lines[i])
                if m2 and not m2[1] and m2[2][0].isdigit() == ordered:
                    items.append(m2[3].strip())
                elif items and lines[i].startswith((" ", "\t")) and lines[i].strip():
                    items[-1] += " " + lines[i].strip()
                else:
                    break
                i += 1
            blocks.append(Block("list", items=items, ordered=ordered))
        else:
            para.append(line)
            i += 1
    flush()
    return blocks


def planned_ranges() -> dict[str, tuple[int, int]]:
    """Each file's word range, from the table in docs/research/README.md."""
    out = {}
    for line in (RESEARCH / "README.md").read_text(encoding="utf-8").splitlines():
        if m := re.match(r"\| `(\d\d_\w+\.md)` \| .*? \| (\d+)-(\d+)", line):
            out[m[1]] = (int(m[2]), int(m[3]))
    return out


# -- XML building blocks ---------------------------------------------------------------------

def el(tag: str, attrs: dict | None = None, children=()) -> etree._Element:
    e = OxmlElement(tag)
    for k, v in (attrs or {}).items():
        e.set(qn(k), str(v))
    for c in children:
        e.append(c)
    return e


def rpr(size=BODY, *, bold=False, italic=False, font=None, color=None, smallcaps=False,
        sup=False):
    kids = []
    if font:
        kids.append(el("w:rFonts", {"w:ascii": font, "w:hAnsi": font, "w:eastAsia": font,
                                    "w:cs": font}))
    if bold:
        kids += [el("w:b"), el("w:bCs")]
    if italic:
        kids += [el("w:i"), el("w:iCs")]
    if smallcaps:
        kids.append(el("w:smallCaps"))
    if color:
        kids.append(el("w:color", {"w:val": color}))
    kids += [el("w:sz", {"w:val": size}), el("w:szCs", {"w:val": size})]
    if sup:
        kids.append(el("w:vertAlign", {"w:val": "superscript"}))
    return el("w:rPr", children=kids)


def text_el(text: str) -> etree._Element:
    t = el("w:t")
    t.text = text
    t.set(XML_SPACE, "preserve")
    return t


def run(text: str, **fmt) -> etree._Element:
    return el("w:r", children=[rpr(**fmt), text_el(text)])


def run_like(text: str, rpr_el) -> etree._Element:
    return el("w:r", children=[copy.deepcopy(rpr_el), text_el(text)] if rpr_el is not None
              else [text_el(text)])


def tab_run(size=BODY) -> etree._Element:
    return el("w:r", children=[rpr(size), el("w:tab")])


def ppr(*, keep_next=False, tabs=None, spacing=None, ind=None, jc=None, num=None):
    kids = []
    if keep_next:
        kids.append(el("w:keepNext"))
    if num:
        kids.append(el("w:numPr", children=[el("w:ilvl", {"w:val": 0}),
                                            el("w:numId", {"w:val": num})]))
    if tabs:
        kids.append(el("w:tabs", children=[el("w:tab", {"w:val": v, "w:pos": pos})
                                           for v, pos in tabs]))
    if spacing:
        kids.append(el("w:spacing", {f"w:{k}": v for k, v in spacing.items()}))
    if ind:
        kids.append(el("w:ind", {f"w:{k}": v for k, v in ind.items()}))
    if jc:
        kids.append(el("w:jc", {"w:val": jc}))
    return el("w:pPr", children=kids)


def paragraph(p_pr=None, children=()) -> etree._Element:
    return el("w:p", children=[p_pr if p_pr is not None else ppr(), *children])


def field_runs(instruction: str, placeholder: str, rpr_el) -> list[etree._Element]:
    def fld(kind):
        return el("w:r", children=[copy.deepcopy(rpr_el), el("w:fldChar", {"w:fldCharType": kind})])
    instr = el("w:instrText")
    instr.text = f" {instruction} "
    instr.set(XML_SPACE, "preserve")
    return [fld("begin"), el("w:r", children=[copy.deepcopy(rpr_el), instr]), fld("separate"),
            run_like(placeholder, rpr_el), fld("end")]


def text_of(e) -> str:
    return "".join(t.text or "" for t in e.iter(qn("w:t")))


# -- maths -----------------------------------------------------------------------------------

class Maths:
    """LaTeX -> MathML (latex2mathml) -> OMML (Office's own stylesheet)."""

    def __init__(self) -> None:
        if not XSL.exists():
            sys.exit(f"MML2OMML.XSL not found at {XSL}; set MML2OMML_XSL")
        self.xslt = etree.XSLT(etree.parse(str(XSL)))

    def omml(self, tex: str, *, display: bool, size: int) -> etree._Element:
        tex = re.sub(r"\\tag\{[^}]*\}", "", tex).strip()
        mml = latex2mathml.converter.convert(tex, display="block" if display else "inline")
        root = self.xslt(etree.fromstring(mml.encode("utf-8"))).getroot()
        omath = root if root.tag == f"{M}oMath" else root.find(f".//{M}oMath")
        omath = copy.deepcopy(omath)
        for r in omath.iter(f"{M}r"):
            self._size(r, size)
        return omath

    @staticmethod
    def _size(r, size: int) -> None:
        w_rpr = r.find(f"{W}rPr")
        if w_rpr is None:
            w_rpr = etree.Element(f"{W}rPr")
            r.insert(1 if r.find(f"{M}rPr") is not None else 0, w_rpr)
        if w_rpr.find(f"{W}rFonts") is None:
            w_rpr.insert(0, etree.Element(f"{W}rFonts", {f"{W}ascii": "Cambria Math",
                                                         f"{W}hAnsi": "Cambria Math"}))
        for tag in ("sz", "szCs"):
            etree.SubElement(w_rpr, f"{W}{tag}", {f"{W}val": str(size)})

    def numbered(self, tex: str, number: str, size: int) -> etree._Element:
        """A display equation with its number flush right, as Word's own `#(n)` builds it."""
        omath = self.omml(tex, display=True, size=size)
        e = etree.Element(f"{M}e")
        e.extend(list(omath))
        for text in ("#", f"({number})"):
            r = etree.SubElement(e, f"{M}r")
            etree.SubElement(etree.SubElement(r, f"{M}rPr"), f"{M}sty", {f"{M}val": "p"})
            self._size(r, size)
            etree.SubElement(r, f"{M}t").text = text
        arr = etree.Element(f"{M}eqArr")
        arr_pr = etree.SubElement(arr, f"{M}eqArrPr")
        etree.SubElement(arr_pr, f"{M}maxDist", {f"{M}val": "1"})
        arr.append(e)
        para = etree.Element(f"{M}oMathPara")
        etree.SubElement(para, f"{M}oMath").append(arr)
        return para


# -- references ------------------------------------------------------------------------------

# Capitalized words in titles that stay capitalized when a title is set in sentence case.
PROPER = {
    "Apache", "Bayesian", "Brier", "Copilot", "Docker", "English", "Gaussian", "Google", "Java",
    "Kubernetes", "Markov", "Monte", "Carlo", "Python", "Scala", "Shapley", "Spark", "Bernoulli",
    "Hadoop", "Bayes", "Kolmogorov", "Smirnov", "Pareto", "Poisson", "Laplace", "European",
    "Microsoft", "Amazon", "Netflix", "Twitter", "Wikipedia", "Jupyter", "Hive", "Delta", "Lake",
    "Parquet", "Arrow", "Pandas", "Dask", "Ray", "Flink", "Kafka", "Airflow", "Kubeflow",
}


def sentence_case(title: str) -> str:
    """Only the first word (and the first after a colon) capitalized, except proper nouns,
    acronyms, and words with internal capitals, which a title keeps as written."""
    out, first = [], True
    for word in title.split(" "):
        core = re.sub(r"^[\"'(\u201c\u2018]+|[\"'),.:;?!\u201d\u2019]+$", "", word)
        if not first and core:
            parts = re.split(r"([-\u2013])", core)
            lowered = "".join(x.lower() if re.fullmatch(r"[A-Z][a-z]+", x) and x not in PROPER
                              else x for x in parts)
            word = word.replace(core, lowered, 1)
        out.append(word)
        first = word.endswith((":", "?", "!", "\u2014"))
    return " ".join(out)


def initials(given: str) -> str:
    out = ""
    for part in re.findall(r"[^\W\d_]+|-", given):
        out += "-" if part == "-" else (" " if out and not out.endswith("-") else "") + part[0] + "."
    return out


def author_list(authors: list[dict]) -> str:
    names = [f"{initials(a['given'])} {a['family']}".strip() if a.get("given") else a["family"]
             for a in authors]
    if len(names) > 6:
        return f"{names[0]} et al."
    if len(names) <= 2:
        return " and ".join(names)
    return ", ".join(names[:-1]) + ", and " + names[-1]


def reference(entry: dict) -> list[tuple[str, bool]]:
    """One IEEE reference as (text, italic) segments."""
    who = author_list(entry["authors"])
    title = sentence_case(entry["title"].strip().rstrip("."))
    pages = entry.get("pages") or ""
    pages = (f"pp. {pages.replace('-', chr(0x2013))}" if "-" in pages
             else f"Art. no. {pages}" if pages else "")
    doi = f"doi: {entry['doi']}."
    kind, container = entry["type"], (entry.get("container") or "").strip()
    lead = f"{who}, " if who else ""
    if kind == "journal-article":
        tail = ", ".join(x for x in (f"vol. {entry['volume']}" if entry.get("volume") else "",
                                     f"no. {entry['issue']}" if entry.get("issue") else "",
                                     pages, str(entry["year"])) if x)
        return [(f"{lead}\u201c{title},\u201d ", False), (container, True), (f", {tail}, {doi}", False)]
    if kind == "proceedings-article":
        tail = ", ".join(x for x in (str(entry["year"]), pages) if x)
        return [(f"{lead}\u201c{title},\u201d in ", False), (container, True),
                (f", {tail}, {doi}", False)]
    if kind == "book":
        series = f" ({container})" if container else ""
        return [(lead, False), (entry["title"].strip(), True),
                (f"{series}. {entry['publisher']}, {entry['year']}, {doi}", False)]
    return [(f"{lead}\u201c{title},\u201d {entry['publisher']}, {entry['year']}, {doi}", False)]


# -- the writer ------------------------------------------------------------------------------

class Writer:
    """Renders blocks into paragraphs and tables, in one list, in reading order."""

    def __init__(self, doc, sect_proto, bibliography: dict) -> None:
        self.doc, self.sect_proto, self.bib = doc, sect_proto, bibliography
        self.maths = Maths()
        self.out: list[etree._Element] = []
        self.order: list[str] = []            # citation keys in order of first citation
        self.columns = 2
        self.first_para = True
        self.words = 0
        self.figures: list[tuple[int, str]] = []
        self.tables: list[tuple[int, str]] = []
        self.equations = 0
        self.missing_figures: list[int] = []
        self.warnings: list[str] = []
        self._bookmarks = itertools.count(9000)

    # layout
    def _section_end(self, cols: int) -> None:
        sect = copy.deepcopy(self.sect_proto)
        sect.find(qn("w:cols")).attrib.clear()
        sect.find(qn("w:cols")).set(qn("w:space"), "708")
        if cols == 2:
            sect.find(qn("w:cols")).set(qn("w:num"), "2")
        last = self.out[-1] if self.out else None
        if last is None or last.tag != qn("w:p") or last.find(qn("w:pPr")) is None \
                or last.find(qn("w:pPr")).find(qn("w:sectPr")) is not None:
            last = paragraph(ppr(spacing={"before": 0, "after": 0}))
            self.out.append(last)
        last.find(qn("w:pPr")).append(sect)

    def wide(self) -> None:
        """Figures, tables, and code span the page: end the two-column run before them."""
        if self.columns == 2:
            self._section_end(2)
            self.columns = 1

    def narrow(self) -> None:
        if self.columns == 1:
            self._section_end(1)
            self.columns = 2

    def bookmark(self, p, name: str) -> None:
        i = next(self._bookmarks)
        p.insert(1, el("w:bookmarkStart", {"w:id": i, "w:name": name}))
        p.append(el("w:bookmarkEnd", {"w:id": i}))

    # inline text
    def inline(self, p, text: str, size=BODY, bold=False, italic=False, **fmt) -> None:
        """Runs for one paragraph's text; `fmt` (font, smallcaps) applies to its words."""
        pos = 0
        for m in TOKEN.finditer(text):
            if m.start() > pos:
                p.append(run(transform(text[pos:m.start()]), size=size, bold=bold, italic=italic,
                             **fmt))
            tok, kind = m.group(), m.lastgroup
            if kind == "code":
                p.append(run(tok[1:-1], size=size - 2, font="Courier New", bold=bold, italic=italic))
            elif kind == "math":
                p.append(self.maths.omml(tok[1:-1], display=False, size=size))
            elif kind == "cite":
                p.append(run(self.cite(tok), size=size, bold=bold, italic=italic, **fmt))
            elif kind == "bold":
                self.inline(p, tok[2:-2], size, True, italic, **fmt)
            else:
                self.inline(p, tok[1:-1], size, bold, True, **fmt)
            pos = m.end()
        if pos < len(text):
            p.append(run(transform(text[pos:]), size=size, bold=bold, italic=italic, **fmt))

    def cite(self, tok: str) -> str:
        nums = []
        for key in re.findall(r"@([^\]\s]+)", tok):
            if key not in self.bib:
                raise KeyError(f"citation key {key!r} is not in references/bibliography.json")
            if key not in self.order:
                self.order.append(key)
            nums.append(self.order.index(key) + 1)
        return "".join(f"[{n}]" for n in nums)

    # blocks
    def heading1(self, label: str | None, title: str, bookmark: str | None = None) -> None:
        self.narrow()
        fmt = {"size": HEAD, "bold": True, "font": HELV, "color": BLUE}
        p = paragraph(ppr(keep_next=True, spacing={"before": 240, "after": 80}),
                      ([run(f"{label}.", **fmt), run(" ", **fmt)] if label else [])
                      + [run(title.upper(), **fmt)])
        if bookmark:
            self.bookmark(p, bookmark)
        self.out.append(p)
        self.first_para = True

    def heading2(self, letter: str, title: str) -> None:
        self.narrow()
        fmt = {"size": HEAD, "bold": True, "italic": True, "font": HELV, "color": GREY}
        self.out.append(paragraph(ppr(keep_next=True, spacing={"before": 200, "after": 60}),
                                  [run(f"{letter}.", **fmt), run(" ", **fmt),
                                   run(transform(title).upper(), **fmt)]))
        self.first_para = True

    def para(self, text: str) -> None:
        self.narrow()
        p = paragraph(ppr(jc="both", ind=None if self.first_para else {"firstLine": 200}))
        self.inline(p, text)
        self.out.append(p)
        self.first_para = False
        self.words += words(text)

    def items(self, items: list[str], ordered: bool) -> None:
        self.narrow()
        for k, item in enumerate(items, 1):
            p = paragraph(ppr(tabs=[("left", 360)], ind={"left": 360, "hanging": 360}, jc="both"),
                          [run(f"{k})" if ordered else "\u2022"), tab_run()])
            self.inline(p, item)
            self.out.append(p)
            self.words += words(item)
        self.first_para = False

    def equation(self, tex: str) -> None:
        self.narrow()
        m = re.search(r"\\tag\{([^}]*)\}", tex)
        if not m:
            raise ValueError(f"display equation without \\tag{{n}}: {tex[:60]}")
        p = paragraph(ppr(spacing={"before": 120, "after": 120}))
        p.append(self.maths.numbered(tex, m[1], BODY))
        self.out.append(p)
        self.equations += 1

    def code(self, text: str) -> None:
        self.wide()
        p = paragraph(ppr(spacing={"before": 80, "after": 120}))
        for k, line in enumerate(text.split("\n")):
            if k:
                p.append(el("w:r", children=[el("w:br")]))
            p.append(run(line, size=CAPTION, font="Courier New"))
        self.out.append(p)

    def figure(self, b: Block) -> None:
        self.wide()
        png = next(iter(sorted((RESEARCH / "figures" / "png").glob(f"Fig {b.number:02d} *.png"))),
                   None)
        if png is None:
            self.missing_figures.append(b.number)
            self.out.append(paragraph(ppr(keep_next=True, jc="center"),
                                      [run(f"[FIGURE {b.number} - EXPORT PENDING]", size=BODY)]))
        else:
            image = Image.from_file(str(png))
            width = TEXT_TWIPS / 1440
            height = width * image.px_height / image.px_width
            if height > MAX_FIGURE_IN:
                width *= MAX_FIGURE_IN / height
            holder = self.doc.add_paragraph()
            holder.add_run().add_picture(str(png), width=Inches(width))
            p = holder._p
            p.insert(0, ppr(keep_next=True, jc="center", spacing={"before": 160, "after": 0}))
            self.out.append(p)
        cap = paragraph(ppr(jc="both", spacing={"before": 80, "after": 200}),
                        [run(f"FIGURE {b.number}.", size=CAPTION, bold=True, font=HELV, color=BLUE),
                         run(" ", size=CAPTION, bold=True, font=HELV)])
        self.inline(cap, b.title, size=CAPTION, bold=True, font=HELV)
        self.bookmark(cap, f"_Fig{b.number:02d}")
        self.out.append(cap)
        self.figures.append((b.number, b.title.rstrip(".")))

    def table(self, rows: list[list[str]], caption: Block | None, part: str | None = None) -> None:
        """A numbered table (caption), a part of one (part, e.g. "Part B ..."), or neither."""
        self.wide()
        if caption is not None:
            label = paragraph(ppr(keep_next=True, jc="center", spacing={"before": 200, "after": 0}),
                              [run(f"TABLE {roman(caption.number)}", size=SMALL, smallcaps=True)])
            self.bookmark(label, f"_Tab{caption.number:02d}")
            title = paragraph(ppr(keep_next=True, jc="center", spacing={"before": 0, "after": 80}))
            self.inline(title, caption.title, size=SMALL, smallcaps=True)
            self.out += [label, title]
            self.tables.append((caption.number, caption.title))
        elif part is None:
            self.warnings.append(f"a table without a caption: {rows[0][:3]}")
        if part is not None:
            sub = paragraph(ppr(keep_next=True, jc="center", spacing={"before": 120, "after": 60}))
            self.inline(sub, part, size=SMALL, italic=True)
            self.out.append(sub)
        self.out.append(self._table(rows))
        note = paragraph(ppr(jc="both", ind={"firstLine": 202}, spacing={"before": 40, "after": 120}))
        if caption is not None and caption.text:
            self.inline(note, caption.text[0].upper() + caption.text[1:].rstrip(".") + ".",
                        size=SMALL)
        self.out.append(note)

    def _table(self, rows: list[list[str]]) -> etree._Element:
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        plain = [[re.sub(r"[*`$]|\[@[^\]]+\]", "", c) for c in r] for r in rows]
        weight = [min(max(len(r[c]) for r in plain), 48) + 6 for c in range(ncols)]
        widths = [int(TEXT_TWIPS * w / sum(weight)) for w in weight]
        line = {"w:val": "single", "w:sz": 12, "w:space": 0, "w:color": "808080"}
        rule = {"w:val": "single", "w:sz": 6, "w:space": 0, "w:color": "000000"}
        tbl = el("w:tbl", children=[
            el("w:tblPr", children=[
                el("w:tblW", {"w:w": 5000, "w:type": "pct"}),
                el("w:tblBorders", children=[el("w:top", line), el("w:bottom", line)]),
                el("w:tblLayout", {"w:type": "fixed"}),
                el("w:tblCellMar", children=[el("w:left", {"w:w": 60, "w:type": "dxa"}),
                                             el("w:right", {"w:w": 60, "w:type": "dxa"})]),
                el("w:tblLook", {"w:val": "0000"}),
            ]),
            el("w:tblGrid", children=[el("w:gridCol", {"w:w": w}) for w in widths]),
        ])
        for ri, row in enumerate(rows):
            tr = el("w:tr", children=[el("w:trPr", children=[el("w:tblHeader")])] if ri == 0 else [])
            for ci, cell in enumerate(row):
                pr = [el("w:tcW", {"w:w": widths[ci], "w:type": "dxa"})]
                if ri == 0:
                    pr.append(el("w:tcBorders", children=[el("w:top", rule), el("w:bottom", rule)]))
                p = paragraph(ppr(jc="center" if ri == 0 else "left",
                                  spacing={"before": 20, "after": 20, "line": 220,
                                           "lineRule": "auto"}))
                self.inline(p, cell, size=SMALL)
                tr.append(el("w:tc", children=[el("w:tcPr", children=pr), p]))
            tbl.append(tr)
        return tbl

    def render(self, blocks: list[Block]) -> None:
        letter, caption, part = 0, None, None
        for k, b in enumerate(blocks):
            # An italic line just above a table names a part of it ("Part B - ...").
            if b.kind == "para" and re.fullmatch(r"\*[^*].*\*", b.text) \
                    and k + 1 < len(blocks) and blocks[k + 1].kind == "table":
                part = b.text[1:-1]
                continue
            if b.kind == "heading":
                letter += 1
                self.heading2(chr(64 + letter),
                              re.sub(r"^(?:\d+(?:\.\d+)*\.?|[A-Z]\.)\s+", "", b.text))
            elif b.kind == "para":
                self.para(b.text)
            elif b.kind == "list":
                self.items(b.items, b.ordered)
            elif b.kind == "equation":
                self.equation(b.text)
            elif b.kind == "code":
                self.code(b.text)
            elif b.kind == "figure":
                self.figure(b)
            elif b.kind == "table_caption":
                caption = b
            elif b.kind == "table":
                self.table(b.rows, caption, part)
                caption, part = None, None

    def references(self) -> None:
        self.heading1(None, "References", bookmark="_Refs")
        for key in self.order:
            p = paragraph(ppr(num=22, jc="both"))
            for text, italic in reference(self.bib[key]):
                p.append(run(text, size=SMALL, italic=italic))
            self.out.append(p)


# -- the template ----------------------------------------------------------------------------

def find(body, predicate, start=0):
    for i, e in enumerate(list(body)[start:], start):
        if predicate(text_of(e).strip()):
            return i, e
    raise LookupError("template paragraph not found")


def set_cell(cell, runs) -> None:
    ps = cell._tc.findall(qn("w:p"))
    for extra in ps[1:]:
        cell._tc.remove(extra)
    for child in list(ps[0]):
        if child.tag != qn("w:pPr"):
            ps[0].remove(child)
    for r in runs:
        ps[0].append(r)


def cell_rpr(table):
    """The header row's run format, unbolded, for the rows below it."""
    for r in table.rows[0]._tr.iter(qn("w:r")):
        rp = r.find(qn("w:rPr"))
        if rp is not None:
            rp = copy.deepcopy(rp)
            for tag in ("w:b", "w:bCs"):
                for x in rp.findall(qn(tag)):
                    rp.remove(x)
            return rp
    return rpr(24)


def fit_rows(table, n: int) -> None:
    """Exactly n data rows below the header: spare template rows removed, missing ones copied."""
    while len(table.rows) - 1 > n:
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) - 1 < n:
        table._tbl.append(copy.deepcopy(table.rows[-1]._tr))


def replace_runs(p, runs, keep_first=False) -> None:
    existing = p.findall(qn("w:r"))
    for r in existing[1 if keep_first else 0:]:
        p.remove(r)
    for r in runs:
        p.append(r)


def split_front(blocks: list[Block]):
    title = next(b.text for b in blocks if b.kind == "heading" and b.level == 1)
    i_abs = next(i for i, b in enumerate(blocks) if b.kind == "heading" and b.text == "Abstract")
    abstract = blocks[i_abs + 1].text
    terms = next(b.text for b in blocks if b.kind == "para" and b.text.startswith("**Index Terms**"))
    terms = re.sub(r"^\*\*Index Terms\*\*\s*-\s*", "", terms).rstrip(".")
    i_sum = next(i for i, b in enumerate(blocks)
                 if b.kind == "heading" and b.text == "Executive Summary")
    return title, abstract, terms, blocks[i_sum + 1:]


def index_terms(terms: str) -> str:
    """IEEE index terms: alphabetical, only the first capitalized, proper nouns kept."""
    out = []
    for k, term in enumerate(sorted((t.strip() for t in terms.split(",")), key=str.lower)):
        first = term.split()[0]
        if k and re.fullmatch(r"[A-Z][a-z]+", first) and first not in PROPER:
            term = term[0].lower() + term[1:]
        out.append(term)
    return ", ".join(out) + "."


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--final", action="store_true", help="refuse pending markers and gaps")
    ap.add_argument("--out", type=Path, default=OUT)
    args = ap.parse_args()

    doc = docx.Document(str(TEMPLATE))
    body = doc.element.body
    bib = json.loads((RESEARCH / "references" / "bibliography.json").read_text(encoding="utf-8"))
    front_tables = doc.tables[:5]       # Index, Figures, Tables, Contributions, Consolidated

    # Template landmarks, found by their text.
    i_intro, _ = find(body, lambda t: t.startswith("I.") and t.endswith("INTRODUCTION"))
    i_fbio, fbio = find(body, lambda t: t.startswith("Dr. M. Rajasekhara Babu received"))
    i_flast, _ = find(body, lambda t: t.startswith("Prof. Rajasekhara Babu is associated"))
    i_sbio, _ = find(body, lambda t: t.startswith("SECOND B. AUTHOR"))
    i_sect = next(i for i, e in enumerate(list(body)) if i > i_sbio
                  and e.find(qn("w:pPr")) is not None
                  and e.find(qn("w:pPr")).find(qn("w:sectPr")) is not None)
    sect_proto = copy.deepcopy(list(body)[i_sect].find(qn("w:pPr")).find(qn("w:sectPr")))
    sect_proto.find(qn("w:type")).set(qn("w:val"), "continuous")

    w = Writer(doc, sect_proto, bib)
    report: dict = {"sections": [], "pending": {}}
    ranges = planned_ranges()

    # Front matter: executive summary in the body, abstract and index terms in the header.
    front, pending = read(RESEARCH / "00_front_matter.md")
    report["pending"]["00_front_matter.md"] = pending
    title, abstract, terms, summary = split_front(front)
    w.heading1(None, "Executive Summary", bookmark="_Exec")
    w.render(summary)
    summary_words = w.words
    front_words = summary_words + words(abstract) + words(terms)
    report["sections"].append({"file": "00_front_matter.md", "topic": "Executive Summary and Abstract",
                               "words": front_words, "abstract_words": words(abstract),
                               "index_terms_words": words(terms),
                               "range": ranges.get("00_front_matter.md")})
    # The Index and the Contributions table take the front matter row by row, with a section's
    # number in its own column, as the course's reference report lays them out: (Sl. No.,
    # topic, bookmark) in the Index; (Sl. No., topic, words) in the Contributions table.
    index_rows = [("", "Abstract", "_Abs"), ("", "Executive Summary", "_Exec")]
    contrib_rows = [("", "Abstract", words(abstract)),
                    ("", "Index Terms (Keywords)", words(terms)),
                    ("", "Executive Summary (with Table 1)", summary_words)]
    da_words = 0

    numbered = sorted(p for p in RESEARCH.glob("[0-2][0-9]_*.md")
                      if re.match(r"(0[1-9]|1\d|2[0-2])_", p.name))
    present = {int(p.name[:2]) for p in numbered}
    for n in range(1, 23):
        if n not in present:
            w.warnings.append(f"Section {n} has no source file")
    for path in numbered:
        blocks, pending = read(path)
        report["pending"][path.name] = pending
        head = blocks[0]
        m = re.match(r"(\d+)\.\s+(.+)", head.text)
        if head.kind != "heading" or not m:
            raise ValueError(f"{path.name} must open with '# N. Title'")
        n, sec_title = int(m[1]), m[2].strip()
        before = w.words
        w.heading1(roman(n), sec_title, bookmark=f"_Sec{n:02d}")
        w.render(blocks[1:])
        count = w.words - before
        report["sections"].append({"file": path.name, "topic": f"{roman(n)}. {sec_title}",
                                   "words": count, "range": ranges.get(path.name)})
        index_rows.append((f"{n}.", sec_title, f"_Sec{n:02d}"))
        contrib_rows.append((f"{n}.", sec_title, count))
        if n == 22:     # the statement belongs with the paper's end matter, after the conclusion
            da, _ = read(RESEARCH / "data_availability.md")
            before = w.words
            w.heading1(None, "Data Availability", bookmark="_DA")
            w.render(da[1:])
            da_words = w.words - before
            report["sections"].append({"file": "data_availability.md",
                                       "topic": "Data Availability", "words": da_words,
                                       "range": None})

    appendix, pending = read(RESEARCH / "23_appendix.md")
    report["pending"]["23_appendix.md"] = pending
    before = w.words
    w.heading1(None, "Appendix", bookmark="_Appx")
    w.render(appendix[1:])
    appendix_words = w.words - before
    report["sections"].append({"file": "23_appendix.md", "topic": "Appendix",
                               "words": appendix_words, "range": ranges.get("23_appendix.md")})
    w.references()
    w.narrow()

    # The template's sample body and sample biographies go; ours goes in their place.
    children = list(body)
    for e in children[i_intro:i_fbio] + children[i_flast + 1:i_sect]:
        body.remove(e)
    for e in w.out:
        fbio.addprevious(e)
    sect_para = children[i_sect]
    bio_text = {"size": SMALL}
    for member in MEMBERS:
        sect_para.addprevious(paragraph(ppr(spacing={"before": 0, "after": 0})))
        sect_para.addprevious(paragraph(ppr(jc="both", spacing={"before": 240}), [
            run(member, size=SMALL, bold=True, font=HELV),
            run(" [BIOGRAPHY AND PHOTOGRAPH - TO BE INSERTED AT FINALIZATION]", **bio_text)]))

    # Cover page.
    for p in children[0].iter(qn("w:p")):
        if text_of(p).strip().startswith("Team ID"):
            last = p.findall(qn("w:r"))[-1]
            p.append(run_like(TEAM_ID, last.find(qn("w:rPr"))))
    for t in body.iter(qn("w:t")):
        if t.text and "<Title>" in t.text:
            t.text = t.text.replace("<Title>", title)
        if t.text and t.text.endswith("On\u201c"):
            t.text = t.text[:-1] + " \u201c"

    # IEEE header.
    _, ptitle = find(body, lambda t: t == "<TITLE OF THE PAPER>")
    for t in ptitle.iter(qn("w:t")):
        t.text = title if "TITLE OF THE PAPER" in (t.text or "") else ""
    _, authors = find(body, lambda t: t.startswith("Dr. M Rajasekhara Babu"))
    arun = authors.findall(qn("w:r"))
    normal = arun[0].find(qn("w:rPr"))
    sup = next(r.find(qn("w:rPr")) for r in arun if r.find(qn("w:rPr")).find(qn("w:vertAlign")) is not None)
    names = [run_like("Dr. M Rajasekhara Babu", normal), run_like("1", sup)]
    for k, member in enumerate(MEMBERS, 2):
        names += [run_like(f", {member}", normal), run_like(str(k), sup)]
    replace_runs(authors, names)
    _, emails = find(body, lambda t: t.startswith("1mrajasekharababu"))
    plain14, sup14 = rpr(CAPTION), rpr(CAPTION, sup=True)
    mail = []
    for k, member in enumerate(MEMBERS, 2):
        mail += [run_like(", ", plain14), run_like(str(k), sup14),
                 run_like(f"{member} VIT E-MAIL", plain14)]
    replace_runs(emails, mail)
    _, abs_p = find(body, lambda t: t.startswith("ABSTRACT"))
    replace_runs(abs_p, [], keep_first=True)
    abs_p.append(run(" ", size=BODY))
    w.inline(abs_p, abstract)
    w.bookmark(abs_p, "_Abs")
    _, terms_p = find(body, lambda t: t.startswith("INDEX TERMS"))
    replace_runs(terms_p, [run(" " + index_terms(terms), size=BODY)], keep_first=True)

    # Index, figures, tables, contributions.
    index, lof, lot, contrib, consolidated = front_tables
    index_rows += [("", "Data Availability", "_DA"), ("", "References", "_Refs"),
                   ("", "Appendix", "_Appx")]
    contrib_rows += [("", "Data Availability", da_words), ("", "References", None),
                     ("", "Appendix", appendix_words)]
    for table, entries in ((index, index_rows), (contrib, contrib_rows)):
        fit_rows(table, len(entries))
        cr = cell_rpr(table)
        for k, (number, topic, value) in enumerate(entries, 1):
            cells = table.rows[k].cells
            set_cell(cells[0], [run_like(number, cr)])
            set_cell(cells[1], [run_like(topic, cr)])
            if table is index:
                set_cell(cells[2], field_runs(f"PAGEREF {value} \\h", "?", cr))
            else:
                set_cell(cells[2], [run_like("[TO BE CONFIRMED]", cr)])
                # References are not counted, as the reference report marks them.
                set_cell(cells[3], [run_like("Excl." if value is None else str(value), cr)])
    for table, entries, label, prefix in ((lof, w.figures, str, "_Fig"),
                                          (lot, sorted(w.tables), roman, "_Tab")):
        fit_rows(table, len(entries))
        cr = cell_rpr(table)
        for k, (number, caption) in enumerate(sorted(entries), 1):
            cells = table.rows[k].cells
            set_cell(cells[0], [run_like(label(number), cr)])
            set_cell(cells[1], [run_like(transform(re.sub(r"[*`]", "", caption)), cr)])
            set_cell(cells[2], field_runs(f"PAGEREF {prefix}{number:02d} \\h", "?", cr))
    total = sum(s["words"] for s in report["sections"])
    cr = cell_rpr(consolidated)
    for k, member in enumerate(MEMBERS, 2):
        cells = consolidated.rows[k].cells
        set_cell(cells[1], [run_like(f"{member} REG. NO.", cr)])
        set_cell(cells[2], [run_like("[TO BE CONFIRMED]", cr)])
    set_cell(consolidated.rows[6].cells[-1], [run_like(str(total), cr)])

    doc.core_properties.title = title
    doc.core_properties.author = f"Team {TEAM_ID}"
    doc.core_properties.subject = "Programming for Data Science (MACSE502) case-study report"

    uncited = sorted(set(bib) - set(w.order))
    report.update({
        "output": str(args.out.relative_to(ROOT)) if args.out.is_relative_to(ROOT) else str(args.out),
        "total_words": total,
        "figures": len(w.figures), "figure_numbers": [n for n, _ in w.figures],
        "missing_figure_files": w.missing_figures,
        "tables": len(w.tables), "table_numbers": sorted(n for n, _ in w.tables),
        "equations": w.equations,
        "references_cited": len(w.order), "bibliography_entries_not_cited": uncited,
        "warnings": w.warnings,
    })
    for s in report["sections"]:
        lo_hi = s["range"]
        s["within_range"] = None if not lo_hi else lo_hi[0] <= s["words"] <= lo_hi[1]
    pending_all = {k: v for k, v in report["pending"].items() if v}
    problems = []
    if pending_all:
        problems.append(f"pending markers in {sorted(pending_all)}")
    if w.missing_figures:
        problems.append(f"figure files missing for {w.missing_figures}")
    if any("no source file" in x for x in w.warnings):
        problems.append("a numbered section has no source file")
    if args.final and problems:
        print("refusing --final:", "; ".join(problems), file=sys.stderr)
        return 1

    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(args.out))
    (args.out.parent / "manuscript_report.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {args.out}: {total} words, {len(w.figures)} figures, {len(w.tables)} tables, "
          f"{w.equations} equations, {len(w.order)} references")
    for s in report["sections"]:
        flag = "" if s["within_range"] in (True, None) else f"  OUTSIDE {s['range']}"
        print(f"  {s['topic'][:60]:60} {s['words']:6}{flag}")
    for x in problems + w.warnings:
        print("  !", x)
    return 0


if __name__ == "__main__":
    sys.exit(main())
